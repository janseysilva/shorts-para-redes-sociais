import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
random.seed(7)
nomes=["Ana Beatriz Souza","Arthur Lima","Bernardo Castro","Camila Rocha","Davi Nogueira","Enzo Farias","Gabriela Moura","Heitor Batista","Isabela Pires","João Pedro Alves","Júlia Teixeira","Laura Cardoso","Lorenzo Dias","Luana Freitas","Lucas Barros","Manuela Cunha","Maria Clara Reis","Matheus Pinto","Miguel Duarte","Nicolas Vieira","Pedro Henrique Sales","Rafael Monteiro","Sofia Andrade","Theo Campos","Valentina Ramos","Vitória Lopes","Yasmin Correia","Samuel Macedo","Helena Brito","Gustavo Tavares"]
nomes.sort()
turmas=["5º Ano A","5º Ano B"]
F="Arial"; hdr=PatternFill("solid",fgColor="217346"); thin=Side(style="thin",color="BFBFBF")
b=Border(left=thin,right=thin,top=thin,bottom=thin)
def header(ws,row,cols):
    for i,c in enumerate(cols,1):
        x=ws.cell(row,i,c); x.font=Font(name=F,size=14,bold=True,color="FFFFFF"); x.fill=hdr
        x.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); x.border=b
rows=[]
faltas_alvo=[2,5,8,25,3,12,30,1,6,22,4,9,0,18,27,7,3,15,21,5,10,2,33,6,11,4,19,8,1,24]
for i,n in enumerate(nomes):
    rows.append((2026001+i,n,turmas[i%2],"Matutino",80,faltas_alvo[i]))
wb=Workbook()
def alunos(ws,fill):
    header(ws,1,["Matrícula","Nome","Turma","Turno","Aulas dadas","Faltas","% Presença"])
    for r,v in enumerate(rows,2):
        for c,val in enumerate(v,1):
            x=ws.cell(r,c,val); x.font=Font(name=F,size=14); x.border=b
            if c!=2: x.alignment=Alignment(horizontal="center")
        g=ws.cell(r,7); g.border=b; g.font=Font(name=F,size=14); g.alignment=Alignment(horizontal="center")
        if fill: g.value=f"=(E{r}-F{r})/E{r}"; g.number_format="0%"
    for col,w in zip("ABCDEFG",[14,30,13,13,13,10,14]): ws.column_dimensions[col].width=w
    ws.row_dimensions[1].height=36; ws.freeze_panes="A2"
ws=wb.active; ws.title="Alunos"; alunos(ws,False)
bs=wb.create_sheet("Busca")
def busca(ws,fill):
    ws["B1"]="Digite a matrícula em B3"; ws["B1"].font=Font(name=F,size=12,italic=True,color="595959")
    for c,t in (("B2","Matrícula"),("C2","Nome"),("D2","Turma")):
        ws[c]=t; ws[c].font=Font(name=F,size=16,bold=True,color="FFFFFF"); ws[c].fill=hdr; ws[c].alignment=Alignment(horizontal="center"); ws[c].border=b
    ws["B3"]=2026015 if fill else None
    for c in ("B3","C3","D3"):
        ws[c].font=Font(name=F,size=16); ws[c].border=b; ws[c].alignment=Alignment(horizontal="center")
    ws["B3"].fill=PatternFill("solid",fgColor="FFF2CC")
    if fill:
        ws["C3"]="=VLOOKUP(B3,Alunos!A:C,2,0)"; ws["D3"]="=VLOOKUP(B3,Alunos!A:C,3,0)"
    ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=18; ws.column_dimensions["C"].width=32; ws.column_dimensions["D"].width=14
    ws.row_dimensions[2].height=28; ws.row_dimensions[3].height=28
busca(bs,False)
gb=wb.create_sheet("Gabarito"); alunos(gb,True)
gb.conditional_formatting.add("G2:G31",CellIsRule(operator="lessThan",formula=["0.75"],fill=PatternFill("solid",fgColor="FFC7CE"),font=Font(color="9C0006")))
gb["I1"]="Busca pronta"; gb["I1"].font=Font(name=F,size=14,bold=True)
gb["I2"]="Matrícula"; gb["J2"]=2026015; gb["I3"]="Nome"; gb["J3"]="=VLOOKUP(J2,A:C,2,0)"; gb["I4"]="Turma"; gb["J4"]="=VLOOKUP(J2,A:C,3,0)"
for c in ("I2","I3","I4","J2","J3","J4"): gb[c].font=Font(name=F,size=14)
gb.column_dimensions["I"].width=14; gb.column_dimensions["J"].width=30
gb["I6"]="Esta aba mostra o resultado final dos 3 shorts. Grave usando as abas Alunos e Busca."
gb["I6"].font=Font(name=F,size=11,italic=True,color="595959")
gb.sheet_properties.tabColor="A6A6A6"
ws["I1"]="Planilha de exemplo com alunos FICTÍCIOS. Nunca grave com dados reais."
ws["I1"].font=Font(name=F,size=11,italic=True,color="C00000")
wb.save("Planilha_Exemplo_Shorts.xlsx")
